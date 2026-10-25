
from __future__ import annotations

import csv
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
PROCESSED = ROOT / "data" / "processed"

WORKBOOK = PROCESSED / "phase2g14_fieldwork_visit_workbook.xlsx"
OUT_ROWS = PROCESSED / "phase2g15_fieldwork_validation.csv"
OUT_SUMMARY = PROCESSED / "phase2g15_fieldwork_validation_summary.csv"
OUT_AUDIT = PROCESSED / "phase2g15_fieldwork_validation_audit.txt"

EXPECTED_CANDIDATES = 255

# Explicit mapping to the actual Phase 2G.14 workbook headers.
FIELDS = {
    "candidate_id": "candidate_record_id",
    "fieldwork_status": "fieldwork_status",
    "visit_date": "field_visit_date",
    "observer": "observer_name_or_id",
    "actual_longitude": "actual_gps_longitude",
    "actual_latitude": "actual_gps_latitude",
    "facility_identity": "facility_identity_result",
    "operating_status": "operating_status",
    "vehicle_access": "vehicle_access_result",
    "road_condition": "observed_road_condition",
    "vehicle_type": "vehicle_type_assessed",
    "evidence_reference": "photo_or_document_reference",
    "field_notes": "field_observation_notes",
    "review_decision": "review_decision",
    "reviewer": "reviewer_name_or_id",
    "review_date": "review_date",
    "review_notes": "review_evidence_notes",
    "physical_access_guardrail": "physical_vehicle_access_confirmed",
    "operational_guardrail": "operational_facility_verified",
    "routing_eligible": "routing_eligible",
}

VISITED_STATUSES = {
    "VISITED", "COMPLETED", "VISIT_COMPLETED", "FIELDWORK_COMPLETED"
}
NOT_VISITED_STATUSES = {
    "NOT_VISITED", "PENDING", "NOT_STARTED", "PLANNED", "AWAITING_FIELDWORK"
}
PENDING_REVIEW_VALUES = {
    "", "PENDING", "PENDING_REVIEW", "NOT_REVIEWED", "AWAITING_REVIEW"
}
VALID_REVIEW_DECISIONS = {
    "APPROVE", "APPROVED", "REJECT", "REJECTED",
    "NEEDS_MORE_EVIDENCE", "REQUIRES_FOLLOW_UP"
}


def clean(value):
    return "" if value is None else str(value).strip()


def key(value):
    return clean(value).strip().upper().replace("-", "_").replace(" ", "_")


def is_blank(value):
    return clean(value) == ""


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if is_blank(value):
        return None

    text = clean(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_number(value):
    if is_blank(value):
        return None
    try:
        return float(clean(value))
    except (ValueError, TypeError):
        return None


def affirmative(value):
    return key(value) in {"YES", "TRUE", "1", "CONFIRMED"}


def normalized_result(value):
    v = key(value)
    if v in {"YES", "TRUE", "1", "CONFIRMED", "OPERATIONAL", "PASS"}:
        return "YES"
    if v in {"NO", "FALSE", "0", "NOT_CONFIRMED", "NOT_OPERATIONAL", "FAIL"}:
        return "NO"
    return "UNKNOWN"


def main():
    if not WORKBOOK.exists():
        print(f"ERROR: Workbook not found: {WORKBOOK}")
        return 2

    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)

    if "All Candidates" not in wb.sheetnames:
        print("ERROR: 'All Candidates' sheet is missing.")
        wb.close()
        return 2

    ws = wb["All Candidates"]
    iterator = ws.iter_rows(values_only=True)
    raw_headers = next(iterator, None)

    if not raw_headers:
        print("ERROR: The candidate worksheet has no header row.")
        wb.close()
        return 2

    headers = [clean(h) for h in raw_headers]
    header_to_index = {h: i for i, h in enumerate(headers) if h}

    missing_columns = [
        (logical, actual)
        for logical, actual in FIELDS.items()
        if actual not in header_to_index
    ]

    records = []
    for excel_row, values in enumerate(iterator, start=2):
        values = list(values)
        if not any(not is_blank(v) for v in values):
            continue
        values += [None] * max(0, len(headers) - len(values))
        records.append((excel_row, values[:len(headers)]))

    def get(row, logical):
        actual = FIELDS[logical]
        idx = header_to_index.get(actual)
        return None if idx is None else row[idx]

    results = []
    issues = []
    ids = []
    counts = Counter()

    for excel_row, row in records:
        cid = clean(get(row, "candidate_id"))
        if cid:
            ids.append(cid)

        row_issues = []

        def flag(code, detail):
            row_issues.append(f"{code}: {detail}")

        if not cid:
            flag("MISSING_CANDIDATE_ID", "Candidate ID is blank.")

        status = key(get(row, "fieldwork_status"))
        visited = status in VISITED_STATUSES
        not_visited = status in NOT_VISITED_STATUSES

        if not visited and not not_visited:
            flag(
                "UNRECOGNIZED_FIELDWORK_STATUS",
                f"Unexpected fieldwork status: {status or '(blank)'}."
            )

        visit_date = parse_date(get(row, "visit_date"))
        observer = get(row, "observer")
        evidence = get(row, "evidence_reference")
        field_notes = get(row, "field_notes")

        lon_raw = get(row, "actual_longitude")
        lat_raw = get(row, "actual_latitude")
        lon = parse_number(lon_raw)
        lat = parse_number(lat_raw)

        gps_valid = (
            lon is not None and lat is not None
            and -180 <= lon <= 180
            and -90 <= lat <= 90
        )

        if visited:
            counts["visited_rows"] += 1

            if is_blank(get(row, "visit_date")):
                flag("VISIT_DATE_MISSING", "Completed visit needs a visit date.")
            elif visit_date is None:
                flag("VISIT_DATE_INVALID", "Visit date is not recognizable.")
            elif visit_date > date.today():
                flag("VISIT_DATE_IN_FUTURE", "Visit date is in the future.")

            if is_blank(observer):
                flag("OBSERVER_MISSING", "Completed visit needs an observer.")

            if is_blank(lon_raw) or is_blank(lat_raw):
                flag("ACTUAL_GPS_MISSING", "Completed visit needs actual field GPS.")
            elif not gps_valid:
                flag("ACTUAL_GPS_INVALID", "Actual GPS coordinates are outside valid ranges or nonnumeric.")

            if is_blank(evidence) and is_blank(field_notes):
                flag("FIELD_EVIDENCE_MISSING", "Add an evidence reference or field notes.")

        elif not not_visited:
            # Do not interpret blank field data as failed fieldwork.
            counts["unrecognized_status_rows"] += 1

        identity = normalized_result(get(row, "facility_identity"))
        operating = normalized_result(get(row, "operating_status"))
        access = normalized_result(get(row, "vehicle_access"))

        road_condition = get(row, "road_condition")
        vehicle_type = get(row, "vehicle_type")

        review_raw = get(row, "review_decision")
        review_value = key(review_raw)
        actual_review = (
            review_value not in PENDING_REVIEW_VALUES
            and review_value != ""
        )

        reviewer = get(row, "reviewer")
        review_date = parse_date(get(row, "review_date"))
        review_notes = get(row, "review_notes")

        if actual_review:
            counts["rows_with_non_pending_review_decisions"] += 1

            if review_value not in VALID_REVIEW_DECISIONS:
                flag(
                    "UNRECOGNIZED_REVIEW_DECISION",
                    f"Unexpected review decision: {review_value}."
                )

            if is_blank(reviewer):
                flag("REVIEWER_MISSING", "A real review decision needs a reviewer.")

            if is_blank(get(row, "review_date")):
                flag("REVIEW_DATE_MISSING", "A real review decision needs a review date.")
            elif review_date is None:
                flag("REVIEW_DATE_INVALID", "Review date is not recognizable.")
            elif review_date > date.today():
                flag("REVIEW_DATE_IN_FUTURE", "Review date is in the future.")

            if is_blank(review_notes) and is_blank(evidence):
                flag(
                    "REVIEW_SUPPORT_MISSING",
                    "A real review decision needs review notes or evidence."
                )

        # These values are safeguards, not evidence that a visit occurred.
        physical_guardrail = normalized_result(
            get(row, "physical_access_guardrail")
        )
        operational_guardrail = normalized_result(
            get(row, "operational_guardrail")
        )

        # Only queue for manual review when the visit and evidence are complete.
        ready_for_manual_review = (
            visited
            and visit_date is not None
            and visit_date <= date.today()
            and not is_blank(observer)
            and gps_valid
            and identity == "YES"
            and operating == "YES"
            and access == "YES"
            and not is_blank(road_condition)
            and not is_blank(vehicle_type)
            and (not is_blank(evidence) or not is_blank(field_notes))
            and actual_review
            and review_value in {"APPROVE", "APPROVED"}
            and not is_blank(reviewer)
            and review_date is not None
            and review_date <= date.today()
            and (not is_blank(review_notes) or not is_blank(evidence))
            and physical_guardrail == "YES"
            and operational_guardrail == "YES"
            and not row_issues
        )

        source_routing = get(row, "routing_eligible")
        if affirmative(source_routing):
            counts["source_routing_yes_rows"] += 1
            flag(
                "ROUTING_ELIGIBILITY_REQUIRES_MANUAL_CONTROL",
                "Source value is YES. This script cannot authorize routing."
            )

        if not_visited:
            disposition = "AWAITING_FIELDWORK"
        elif ready_for_manual_review:
            disposition = "READY_FOR_MANUAL_ROUTING_REVIEW"
        elif row_issues:
            disposition = "REVIEW_REQUIRED"
        elif visited:
            disposition = "INCOMPLETE_EVIDENCE_OR_REVIEW"
        else:
            disposition = "STATUS_REQUIRES_CLARIFICATION"

        results.append({
            "excel_row": excel_row,
            "candidate_record_id": cid,
            "fieldwork_status": clean(get(row, "fieldwork_status")),
            "field_visit_date": "" if visit_date is None else visit_date.isoformat(),
            "observer_name_or_id": clean(observer),
            "actual_gps_longitude": "" if lon is None else lon,
            "actual_gps_latitude": "" if lat is None else lat,
            "actual_gps_valid": "YES" if gps_valid else "NO",
            "facility_identity_result": identity,
            "operating_status": operating,
            "vehicle_access_result": access,
            "observed_road_condition_present": "YES" if not is_blank(road_condition) else "NO",
            "vehicle_type_assessed_present": "YES" if not is_blank(vehicle_type) else "NO",
            "evidence_reference_present": "YES" if not is_blank(evidence) else "NO",
            "field_notes_present": "YES" if not is_blank(field_notes) else "NO",
            "review_decision": clean(review_raw),
            "review_is_non_pending": "YES" if actual_review else "NO",
            "reviewer_present": "YES" if not is_blank(reviewer) else "NO",
            "review_date": "" if review_date is None else review_date.isoformat(),
            "physical_vehicle_access_confirmed": physical_guardrail,
            "operational_facility_verified": operational_guardrail,
            "source_routing_eligible": clean(source_routing),
            "validation_issue_count": len(row_issues),
            "validation_issues": " | ".join(row_issues),
            "validation_disposition": disposition,
            "routing_approval_granted_by_script": "NO",
        })

        for issue in row_issues:
            issues.append({
                "excel_row": excel_row,
                "candidate_record_id": cid,
                "issue": issue,
            })

        counts[disposition] += 1
        if row_issues:
            counts["rows_with_validation_issues"] += 1

    wb.close()

    duplicate_ids = sorted(
        cid for cid, n in Counter(ids).items() if n > 1
    )
    for duplicate in duplicate_ids:
        issues.append({
            "excel_row": "",
            "candidate_record_id": duplicate,
            "issue": "DUPLICATE_CANDIDATE_ID: ID occurs more than once.",
        })

    missing_id_rows = sum(not r["candidate_record_id"] for r in results)

    output_fields = list(results[0].keys()) if results else [
        "excel_row", "candidate_record_id", "validation_disposition"
    ]
    with OUT_ROWS.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=output_fields)
        writer.writeheader()
        writer.writerows(results)

    summary_rows = [
        ("workbook", str(WORKBOOK)),
        ("worksheet", "All Candidates"),
        ("candidate_rows_read", len(results)),
        ("expected_candidate_rows", EXPECTED_CANDIDATES),
        ("candidate_row_count_matches_expected", len(results) == EXPECTED_CANDIDATES),
        ("unique_candidate_ids", len(set(ids))),
        ("missing_candidate_id_rows", missing_id_rows),
        ("duplicate_candidate_id_values", len(duplicate_ids)),
        ("visited_rows", counts["visited_rows"]),
        ("awaiting_fieldwork_rows", counts["AWAITING_FIELDWORK"]),
        ("rows_with_non_pending_review_decisions", counts["rows_with_non_pending_review_decisions"]),
        ("source_routing_eligible_yes_rows", counts["source_routing_yes_rows"]),
        ("ready_for_manual_routing_review_rows", counts["READY_FOR_MANUAL_ROUTING_REVIEW"]),
        ("review_required_rows", counts["REVIEW_REQUIRED"]),
        ("incomplete_evidence_or_review_rows", counts["INCOMPLETE_EVIDENCE_OR_REVIEW"]),
        ("rows_with_validation_issues", counts["rows_with_validation_issues"]),
        ("total_validation_issues", len(issues)),
        ("missing_expected_columns", "; ".join(actual for _, actual in missing_columns)),
        ("routing_approval_granted_by_script", "NO"),
    ]

    with OUT_SUMMARY.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["metric", "value"])
        writer.writerows(summary_rows)
        writer.writerow([])
        writer.writerow(["issue_type", "count"])
        for issue_type, count in sorted(
            Counter(i["issue"].split(":", 1)[0] for i in issues).items()
        ):
            writer.writerow([issue_type, count])

    lines = [
        "PNIS PHASE 2G.15 — FIELDWORK RESULTS VALIDATION",
        "=" * 72,
        f"Generated: {datetime.now().isoformat(timespec='seconds')}",
        "Status: " + (
            "SCHEMA ERROR" if missing_columns
            else "COMPLETED WITH VALIDATION RESULTS"
        ),
        "",
        "INPUT",
        str(WORKBOOK),
        "Worksheet: All Candidates",
        "",
        "RESULTS",
        f"Candidate rows read: {len(results)}",
        f"Expected candidate rows: {EXPECTED_CANDIDATES}",
        f"Unique candidate IDs: {len(set(ids))}",
        f"Missing candidate IDs: {missing_id_rows}",
        f"Duplicate candidate ID values: {len(duplicate_ids)}",
        f"Visited rows: {counts['visited_rows']}",
        f"Awaiting fieldwork rows: {counts['AWAITING_FIELDWORK']}",
        f"Rows with actual non-pending review decisions: {counts['rows_with_non_pending_review_decisions']}",
        f"Rows ready for manual routing review: {counts['READY_FOR_MANUAL_ROUTING_REVIEW']}",
        f"Rows requiring review: {counts['REVIEW_REQUIRED']}",
        f"Rows with incomplete evidence/review: {counts['INCOMPLETE_EVIDENCE_OR_REVIEW']}",
        f"Rows with validation issues: {counts['rows_with_validation_issues']}",
        f"Total validation issues: {len(issues)}",
        f"Source routing-eligible YES rows: {counts['source_routing_yes_rows']}",
        "",
        "COLUMN MAPPING",
    ]

    for logical, actual in FIELDS.items():
        lines.append(
            f"{logical}: {actual} - "
            f"{'FOUND' if actual in header_to_index else 'MISSING'}"
        )

    lines.extend([
        "",
        "MISSING EXPECTED COLUMNS",
        *(
    [actual for _, actual in missing_columns]
    if missing_columns else ["None"]
),
        "",
        "INTERPRETATION",
        "- PENDING_REVIEW is not treated as a completed review.",
        "- NOT_VISITED candidates are classified as AWAITING_FIELDWORK.",
        "- Blank field observations are not treated as failed inspections.",
        "- Source guardrail values are reported, not overwritten.",
        "- No candidate is approved for routing by this script.",
        "- READY_FOR_MANUAL_ROUTING_REVIEW is only a review queue.",
        "- Physical access requires actual field evidence; road proximity is insufficient.",
        "",
        "OUTPUTS",
        str(OUT_ROWS),
        str(OUT_SUMMARY),
        str(OUT_AUDIT),
    ])
    OUT_AUDIT.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print("=" * 72)
    print("PNIS PHASE 2G.15 — FIELDWORK RESULTS VALIDATION")
    print("=" * 72)
    print(f"Candidate rows read: {len(results)}")
    print(f"Unique candidate IDs: {len(set(ids))}")
    print(f"Duplicate IDs: {len(duplicate_ids)}")
    print(f"Visited rows: {counts['visited_rows']}")
    print(f"Awaiting fieldwork: {counts['AWAITING_FIELDWORK']}")
    print(f"Non-pending review decisions: {counts['rows_with_non_pending_review_decisions']}")
    print(f"Rows with validation issues: {counts['rows_with_validation_issues']}")
    print(f"Total validation issues: {len(issues)}")
    print(f"Missing expected columns: {len(missing_columns)}")
    print()
    print("OUTPUT FILES")
    print(OUT_ROWS)
    print(OUT_SUMMARY)
    print(OUT_AUDIT)

    if missing_columns:
        print("\nSTATUS: SCHEMA ERROR — one or more expected columns are missing.")
        return 2

    if len(results) != EXPECTED_CANDIDATES:
        print("\nSTATUS: REVIEW — candidate row count differs from 255.")
        return 1

    if missing_id_rows or duplicate_ids:
        print("\nSTATUS: REVIEW — candidate ID integrity checks need attention.")
        return 1

    print("\nSTATUS: COMPLETED — schema and candidate ID checks passed.")
    print("No routing approvals were granted.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
