
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[3]
path = root / "data" / "processed" / "phase2g14_fieldwork_visit_workbook.xlsx"

if not path.exists():
    raise FileNotFoundError(path)

wb = load_workbook(path, read_only=True, data_only=True)

print("=" * 78)
print("PHASE 2G.15 — WORKBOOK HEADER INSPECTION (READ-ONLY)")
print("=" * 78)
print("Workbook:", path)
print("Sheets:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n--- SHEET: {sheet_name} ---")
    print("Dimensions:", ws.max_row, "rows x", ws.max_column, "columns")

    if ws.max_row < 1:
        continue

    headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    for i, header in enumerate(headers, start=1):
        print(f"{i:02d}. {header!r}")

    if sheet_name == "All Candidates":
        print("\nFirst three candidate rows:")
        for row_number, row in enumerate(
            ws.iter_rows(min_row=2, max_row=min(4, ws.max_row), values_only=True),
            start=2
        ):
            print(f"\nExcel row {row_number}:")
            for header, value in zip(headers, row):
                if value is not None and str(value).strip():
                    print(f"  {header!r} = {value!r}")

wb.close()
print("\nInspection complete. No workbook data was changed.")
