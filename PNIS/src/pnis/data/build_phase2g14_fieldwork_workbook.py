
from __future__ import annotations

import csv
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "data" / "processed"

TRACKER_FILE = DATA / "phase2g13_fieldwork_execution_tracker.csv"
READINESS_FILE = DATA / "phase2g12_facility_access_readiness_register.csv"

OUTPUT_FILE = DATA / "phase2g14_fieldwork_visit_workbook.xlsx"
AUDIT_FILE = DATA / "phase2g14_fieldwork_workbook_audit.txt"

ID_FIELD = "candidate_record_id"

FIELD_COLUMNS = [
    "fieldwork_priority_order",
    ID_FIELD,
    "access_risk",
    "nearest_mapped_road_distance_m",
    "nearest_road_fclass",
    "planned_longitude",
    "planned_latitude",
    "fieldwork_status",
    "field_visit_date",
    "observer_name_or_id",
    "actual_gps_longitude",
    "actual_gps_latitude",
    "facility_identity_result",
    "operating_status",
    "vehicle_access_result",
    "observed_road_condition",
    "vehicle_type_assessed",
    "photo_or_document_reference",
    "field_observation_notes",
    "review_decision",
    "reviewer_name_or_id",
    "review_date",
    "review_evidence_notes",
    "physical_vehicle_access_confirmed",
    "operational_facility_verified",
    "routing_eligible",
]

def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"Missing input: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        if not reader.fieldnames:
            raise ValueError(f"CSV has no header: {path.name}")
        rows = list(reader)
    if ID_FIELD not in reader.fieldnames:
        raise ValueError(f"{ID_FIELD} missing from {path.name}")
    return rows

def unique_map(rows: list[dict[str, str]], label: str) -> dict[str, dict[str, str]]:
    result = {}
    for row in rows:
        candidate_id = (row.get(ID_FIELD) or "").strip()
        if not candidate_id:
            raise ValueError(f"Blank candidate ID in {label}")
        if candidate_id in result:
            raise ValueError(f"Duplicate candidate ID {candidate_id} in {label}")
        result[candidate_id] = row
    return result

def risk_value(row: dict[str, str]) -> str:
    value = (row.get("phase2g5_access_risk_screen_v2") or "").strip()
    return value or "UNCLASSIFIED"

def style_sheet(ws, freeze="A2", filter_enabled=True):
    ws.freeze_panes = freeze
    if filter_enabled and ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234A70")
        cell.alignment = Alignment(
            vertical="center", horizontal="center", wrap_text=True
        )
    ws.row_dimensions[1].height = 36

    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max(
            (len(str(c.value or "")) for c in column_cells[:100]),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(max_length + 2, 13), 32)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

def add_dropdown(ws, header_name, options):
    headers = [cell.value for cell in ws[1]]
    if header_name not in headers or ws.max_row < 2:
        return
    column_index = headers.index(header_name) + 1
    letter = get_column_letter(column_index)
    formula = '"' + ",".join(options) + '"'
    validation = DataValidation(
        type="list", formula1=formula, allow_blank=True
    )
    validation.errorTitle = "Invalid value"
    validation.error = "Select a value from the dropdown list."
    validation.promptTitle = header_name
    validation.prompt = "Choose a listed value, or leave blank if not yet known."
    ws.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{ws.max_row}")

def populate_sheet(ws, rows, columns):
    ws.append(columns)
    for row in rows:
        ws.append([row.get(column, "") for column in columns])
    style_sheet(ws)

def main():
    print("=" * 72)
    print("PNIS PHASE 2G.14 - FIELDWORK VISIT WORKBOOK")
    print("=" * 72)

    tracker_rows = read_csv(TRACKER_FILE)
    readiness_rows = read_csv(READINESS_FILE)

    tracker = unique_map(tracker_rows, "Phase 2G.13 tracker")
    readiness = unique_map(readiness_rows, "Phase 2G.12 readiness register")

    issues = []
    if set(tracker) != set(readiness):
        issues.append("Candidate IDs do not match between tracker and readiness register.")
        raise ValueError(issues[-1])

    enriched = []
    for row in tracker_rows:
        candidate_id = (row.get(ID_FIELD) or "").strip()
        ready_row = readiness[candidate_id]
        combined = dict(row)
        combined["access_risk"] = risk_value(ready_row)
        enriched.append(combined)

    # Preserve the tracker order and select high-risk candidates by the
    # explicit risk classification, not by assumptions about row position.
    high_risk = [
        row for row in enriched if row["access_risk"] == "HIGH_REVIEW"
    ]

    if len(enriched) != 255:
        issues.append(f"Expected 255 tracker records; found {len(enriched)}.")
    if len({r[ID_FIELD] for r in enriched}) != len(enriched):
        issues.append("Duplicate IDs detected in workbook source.")
    if any(r.get("routing_eligible") != "NO" for r in enriched):
        issues.append("A source candidate is unexpectedly marked routing-eligible.")
    if len(high_risk) != 25:
        issues.append(f"Expected 25 high-risk candidates; found {len(high_risk)}.")

    if issues:
        raise ValueError("\n".join(issues))

    wb = Workbook()
    instructions = wb.active
    instructions.title = "Read Me"

    instruction_rows = [
        ["PNIS Phase 2G.14 — Fieldwork Visit Workbook"],
        ["Purpose", "Record actual observations and evidence for candidate locations."],
        ["All Candidates", "Contains all 255 candidates in Phase 2G.13 tracker order."],
        ["High Risk 25", "Filtered list of candidates classified HIGH_REVIEW."],
        ["How to record", "Enter only observations actually collected during a field visit."],
        ["Date format", "Use YYYY-MM-DD for field visit and review dates."],
        ["GPS", "Record actual GPS coordinates at the observed location, not planned coordinates."],
        ["Evidence", "Use a photo/document filename or reference that your team can retrieve."],
        ["Unknown values", "Leave blank if not observed; do not guess."],
        ["Review", "A reviewer should document the evidence and decision after the visit."],
        ["Safety rule", "Road proximity alone does not prove physical vehicle access."],
        ["Routing safeguard", "Do not change routing_eligible to YES without documented evidence review and approval."],
        ["Preservation", "This workbook is separate from the Phase 2G.13 tracker and source files."],
        ["Generated UTC", datetime.now(timezone.utc).isoformat()],
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 25
    instructions.column_dimensions["B"].width = 105
    instructions.freeze_panes = "A2"
    instructions.sheet_view.showGridLines = False
    instructions["A1"].font = Font(bold=True, size=15, color="FFFFFF")
    instructions["A1"].fill = PatternFill("solid", fgColor="234A70")
    instructions["B1"].fill = PatternFill("solid", fgColor="234A70")
    for row in instructions.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    instructions.row_dimensions[1].height = 30

    all_ws = wb.create_sheet("All Candidates")
    populate_sheet(all_ws, enriched, FIELD_COLUMNS)

    high_ws = wb.create_sheet("High Risk 25")
    populate_sheet(high_ws, high_risk, FIELD_COLUMNS)

    # Controlled dropdown values help reduce inconsistent data entry.
    for ws in (all_ws, high_ws):
        add_dropdown(ws, "fieldwork_status", [
            "NOT_VISITED", "VISIT_SCHEDULED",
            "VISIT_COMPLETED", "FOLLOW_UP_REQUIRED"
        ])
        add_dropdown(ws, "facility_identity_result", [
            "CONFIRMED", "NOT_FOUND", "MISMATCH", "UNCERTAIN"
        ])
        add_dropdown(ws, "operating_status", [
            "OPERATIONAL", "NOT_OPERATIONAL", "SEASONAL",
            "TEMPORARILY_CLOSED", "UNKNOWN"
        ])
        add_dropdown(ws, "vehicle_access_result", [
            "ACCESS_OBSERVED", "NO_VEHICLE_ACCESS_OBSERVED",
            "ACCESS_RESTRICTED", "NOT_ASSESSED"
        ])
        add_dropdown(ws, "review_decision", [
            "PENDING_REVIEW", "ACCEPTED", "REJECTED",
            "FOLLOW_UP_REQUIRED"
        ])
        add_dropdown(ws, "physical_vehicle_access_confirmed", ["YES", "NO"])
        add_dropdown(ws, "operational_facility_verified", ["YES", "NO"])
        add_dropdown(ws, "routing_eligible", ["YES", "NO"])

    # Highlight fields that require field evidence or reviewer input.
    input_headers = {
        "field_visit_date", "observer_name_or_id",
        "actual_gps_longitude", "actual_gps_latitude",
        "facility_identity_result", "operating_status",
        "vehicle_access_result", "observed_road_condition",
        "vehicle_type_assessed", "photo_or_document_reference",
        "field_observation_notes", "reviewer_name_or_id",
        "review_date", "review_evidence_notes",
    }
    for ws in (all_ws, high_ws):
        headers = [c.value for c in ws[1]]
        for header in input_headers:
            if header in headers:
                col = headers.index(header) + 1
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, col).fill = PatternFill("solid", fgColor="FFF2CC")

    wb.save(OUTPUT_FILE)

    # Reopen to verify workbook integrity and row counts.
    check_wb = load_workbook(OUTPUT_FILE, read_only=True, data_only=True)
    check_all = check_wb["All Candidates"].max_row - 1
    check_high = check_wb["High Risk 25"].max_row - 1
    sheet_names = check_wb.sheetnames
    check_wb.close()

    checks = [
        ("Workbook file exists", OUTPUT_FILE.exists()),
        ("All Candidates contains 255 records", check_all == 255),
        ("High Risk 25 contains 25 records", check_high == 25),
        ("Workbook contains instructions and both candidate sheets",
         sheet_names == ["Read Me", "All Candidates", "High Risk 25"]),
        ("Source tracker remains routing-ineligible",
         all(r.get("routing_eligible") == "NO" for r in enriched)),
    ]
    failures = [name for name, passed in checks if not passed]

    audit = [
        "PNIS PHASE 2G.14 - FIELDWORK WORKBOOK AUDIT",
        "=" * 72,
        f"Generated UTC: {datetime.now(timezone.utc).isoformat()}",
        f"Source tracker rows: {len(tracker_rows)}",
        f"High-risk candidates: {len(high_risk)}",
        f"Workbook all-candidate rows: {check_all}",
        f"Workbook high-risk rows: {check_high}",
        f"Workbook sheets: {', '.join(sheet_names)}",
        "",
        "CHECK RESULTS",
        "-" * 72,
    ]
    audit.extend(
        f'{"PASS" if passed else "FAIL"}: {name}'
        for name, passed in checks
    )
    audit.extend([
        "",
        "OUTPUT",
        "-" * 72,
        str(OUTPUT_FILE),
        "",
        "LIMITATIONS",
        "-" * 72,
        "Workbook preparation is not field verification.",
        "No candidate is approved for routing by this script.",
        "Dropdown selections do not replace evidence review.",
        "All changes should be based on actual observations and documented review.",
        "",
        f"Checks passed: {sum(p for _, p in checks)}/{len(checks)}",
        f"Checks failed: {len(failures)}/{len(checks)}",
        f"Status: {'PASS' if not failures else 'REVIEW_REQUIRED'}",
    ])
    AUDIT_FILE.write_text("\n".join(audit) + "\n", encoding="utf-8")

    print(f"Workbook: {OUTPUT_FILE}")
    print(f"Audit:    {AUDIT_FILE}")
    print(f"All candidates: {check_all}")
    print(f"High-risk candidates: {check_high}")
    print(f"Checks passed: {sum(p for _, p in checks)}/{len(checks)}")
    print(f"Checks failed: {len(failures)}/{len(checks)}")
    print(f"STATUS: {'PASS' if not failures else 'REVIEW_REQUIRED'}")

    if failures:
        raise SystemExit(1)

if __name__ == "__main__":
    main()
