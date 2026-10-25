
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
PROCESSED = ROOT / "data" / "processed"

INPUT = PROCESSED / "phase2g18_risk_prioritized_worklist.csv"
WORKBOOK_OUT = PROCESSED / "phase2g19_fieldwork_execution_package.xlsx"
CSV_OUT = PROCESSED / "phase2g19_field_observation_template.csv"
AUDIT_OUT = PROCESSED / "phase2g19_fieldwork_execution_audit.txt"

EXPECTED_ROWS = 255
REQUIRED_COLUMNS = [
    "candidate_record_id",
    "risk_class",
    "fieldwork_priority_order",
]

# These fields are added as blank observation fields. They are not
# assertions about candidate status and must be completed in the field.
OBSERVATION_FIELDS = [
    "field_visit_date",
    "field_team_or_reviewer",
    "gps_location_checked",
    "observed_latitude",
    "observed_longitude",
    "facility_exists",
    "facility_name_observed",
    "facility_type_observed",
    "operational_facility_verified",
    "physical_vehicle_access_confirmed",
    "vehicle_type_assessed",
    "road_surface_observed",
    "road_condition_observed",
    "seasonal_access_concern",
    "safe_vehicle_turnaround",
    "nearest_road_connection_verified",
    "photo_or_evidence_reference",
    "local_contact_or_source",
    "field_observations",
    "follow_up_required",
    "follow_up_details",
    "review_decision",
    "reviewer_notes",
    "routing_approval_recommendation",
    "routing_approval_granted",
]

YES_NO = '"YES,NO,UNKNOWN"'
FACILITY_OPTIONS = '"YES,NO,UNCERTAIN"'
DECISION_OPTIONS = '"PENDING,VERIFIED,REJECTED,NEEDS_FOLLOW_UP"'
ROAD_SURFACE_OPTIONS = (
    '"PAVED,GRAVEL,EARTH,TRACK,FOOTPATH,OTHER,UNKNOWN"'
)
VEHICLE_OPTIONS = '"MOTORCYCLE,PICKUP,VAN,TRUCK,OTHER,NOT_ASSESSED"'
APPROVAL_OPTIONS = '"NO,UNDER_REVIEW,APPROVED"'

def stop(message: str) -> None:
    raise SystemExit(f"ERROR: {message}")

def csv_rows(path: Path) -> list[dict]:
    if not path.exists():
        stop(f"Required input file not found: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def validate_input(rows: list[dict]) -> list[str]:
    if not rows:
        stop("The Phase 2G.18 worklist is empty.")

    columns = list(rows[0].keys())
    missing = [c for c in REQUIRED_COLUMNS if c not in columns]
    if missing:
        stop(f"Required columns missing from input: {missing}")

    if len(rows) != EXPECTED_ROWS:
        stop(
            f"Expected {EXPECTED_ROWS} rows but found {len(rows)}. "
            "No outputs have been written."
        )

    ids = [r["candidate_record_id"].strip() for r in rows]
    if any(not x for x in ids):
        stop("Blank candidate_record_id found.")
    if len(set(ids)) != len(ids):
        stop("Duplicate candidate_record_id values found.")

    for row in rows:
        try:
            int(row["fieldwork_priority_order"])
        except (ValueError, TypeError):
            stop(
                "Invalid fieldwork_priority_order for "
                + row["candidate_record_id"]
            )

    orders = [int(r["fieldwork_priority_order"]) for r in rows]
    if len(set(orders)) != len(orders):
        stop("Duplicate fieldwork priority order values found.")

    return columns

def style_sheet(ws, freeze="A2", filter_enabled=True):
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    if filter_enabled and ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = Border(bottom=thin_gray)

    ws.row_dimensions[1].height = 34

    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        header = str(column_cells[0].value or "")
        max_len = len(header)

        for cell in column_cells[1:]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 45))
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 34)

    for row_num in range(2, ws.max_row + 1):
        ws.row_dimensions[row_num].height = 30

def add_dropdown(ws, header_names, field_name, options):
    if field_name not in header_names:
        return
    col = header_names.index(field_name) + 1
    letter = get_column_letter(col)
    dv = DataValidation(
        type="list", formula1=options, allow_blank=True
    )
    dv.error = "Select a value from the dropdown list."
    dv.errorTitle = "Invalid value"
    dv.prompt = "Select an observed value. Leave blank until assessed."
    dv.promptTitle = "Field observation"
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")

def write_data_sheet(wb, title, source_columns, rows):
    ws = wb.create_sheet(title)
    headers = source_columns + OBSERVATION_FIELDS
    ws.append(headers)

    for source_row in rows:
        values = [source_row.get(col, "") for col in source_columns]
        ws.append(values + [""] * len(OBSERVATION_FIELDS))

    header_names = headers
    style_sheet(ws)

    add_dropdown(ws, header_names, "gps_location_checked", YES_NO)
    add_dropdown(ws, header_names, "facility_exists", FACILITY_OPTIONS)
    add_dropdown(
        ws, header_names, "operational_facility_verified", YES_NO
    )
    add_dropdown(
        ws, header_names, "physical_vehicle_access_confirmed", YES_NO
    )
    add_dropdown(ws, header_names, "vehicle_type_assessed", VEHICLE_OPTIONS)
    add_dropdown(ws, header_names, "road_surface_observed", ROAD_SURFACE_OPTIONS)
    add_dropdown(ws, header_names, "seasonal_access_concern", YES_NO)
    add_dropdown(ws, header_names, "safe_vehicle_turnaround", YES_NO)
    add_dropdown(
        ws, header_names, "nearest_road_connection_verified", YES_NO
    )
    add_dropdown(ws, header_names, "follow_up_required", YES_NO)
    add_dropdown(ws, header_names, "review_decision", DECISION_OPTIONS)
    add_dropdown(
        ws,
        header_names,
        "routing_approval_recommendation",
        '"NO,UNDER_REVIEW,APPROVED"',
    )
    # The actual approval field is deliberately restricted to NO at
    # package creation. Reviewers must not use this template to grant it.
    add_dropdown(ws, header_names, "routing_approval_granted", '"NO"')

    # Distinguish new field observations from imported desk-review data.
    for field in OBSERVATION_FIELDS:
        if field in header_names:
            col = header_names.index(field) + 1
            for row_num in range(2, ws.max_row + 1):
                ws.cell(row_num, col).fill = PatternFill(
                    "solid", fgColor="FFF2CC"
                )

    return ws

def main():
    print("=" * 72)
    print("PNIS PHASE 2G.19 - FIELDWORK EXECUTION PACKAGE")
    print("=" * 72)

    PROCESSED.mkdir(parents=True, exist_ok=True)

    rows = csv_rows(INPUT)
    source_columns = validate_input(rows)

    # Preserve the priority sequence established in Phase 2G.10/2G.18.
    rows = sorted(rows, key=lambda r: int(r["fieldwork_priority_order"]))

    high_risk = [r for r in rows if r["risk_class"].strip() == "HIGH_REVIEW"]
    elevated = [
        r for r in rows if r["risk_class"].strip() == "ELEVATED_REVIEW"
    ]
    standard = [
        r for r in rows if r["risk_class"].strip() == "STANDARD_REVIEW"
    ]

    if len(high_risk) != 25:
        stop(f"Expected 25 HIGH_REVIEW candidates; found {len(high_risk)}.")
    if len(elevated) != 13 or len(standard) != 217:
        stop(
            "Risk counts differ from validated Phase 2G.18 counts. "
            "No outputs have been written."
        )

    wb = Workbook()
    instructions = wb.active
    instructions.title = "Read Me"

    readme_rows = [
        ["PNIS PHASE 2G.19 - FIELDWORK EXECUTION PACKAGE"],
        ["Purpose", "Record field observations for the 255 screened candidates."],
        ["Source", str(INPUT)],
        ["Candidate count", str(len(rows))],
        ["Risk counts", "HIGH_REVIEW=25; ELEVATED_REVIEW=13; STANDARD_REVIEW=217"],
        ["Priority order", "Existing Phase 2G.18 fieldwork_priority_order is preserved."],
        ["Important", "Risk classes are desk-based flags, not verified field findings."],
        ["Before visiting", "Confirm safe access, local permissions, and field-team safety."],
        ["Evidence", "Record date, reviewer, GPS checks, observations, and evidence references."],
        ["Unknown values", "Use UNKNOWN or UNCERTAIN where available; do not guess."],
        ["Blank fields", "Yellow cells are new observation fields; leave blank until assessed."],
        ["Routing guardrail", "This package does not approve any candidate for routing."],
        ["Approval field", "routing_approval_granted is restricted to NO in this template."],
        ["High-risk sheet", "Use High Risk First for the initial review pass; all candidates remain listed."],
        ["Data integrity", "Do not overwrite source files; retain evidence and review changes separately."],
    ]
    for row in readme_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 25
    instructions.column_dimensions["B"].width = 105
    instructions.freeze_panes = "A2"
    instructions.sheet_view.showGridLines = False
    instructions["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    instructions["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    instructions["B1"].fill = PatternFill("solid", fgColor="1F4E78")
    for row_num in range(2, instructions.max_row + 1):
        instructions.cell(row_num, 1).font = Font(bold=True)
        instructions.cell(row_num, 1).alignment = Alignment(
            vertical="top", wrap_text=True
        )
        instructions.cell(row_num, 2).alignment = Alignment(
            vertical="top", wrap_text=True
        )
        instructions.row_dimensions[row_num].height = 32

    all_ws = write_data_sheet(wb, "All Candidates", source_columns, rows)
    high_ws = write_data_sheet(
        wb, "High Risk First", source_columns, high_risk
    )

    # Add a dedicated, concise observation sheet for field teams.
    form = wb.create_sheet("Field Observation Form")
    form_headers = [
        "candidate_record_id",
        "fieldwork_priority_order",
        "risk_class",
        "risk_reason",
        "longitude",
        "latitude",
        "nearest_mapped_road_distance_m",
        "nearest_road_fclass",
    ] + OBSERVATION_FIELDS
    form.append(form_headers)

    for row in rows:
        form.append(
            [row.get(col, "") for col in form_headers[:8]]
            + [""] * len(OBSERVATION_FIELDS)
        )

    form_headers_actual = form_headers
    style_sheet(form)
    add_dropdown(form, form_headers_actual, "gps_location_checked", YES_NO)
    add_dropdown(form, form_headers_actual, "facility_exists", FACILITY_OPTIONS)
    add_dropdown(
        form, form_headers_actual, "operational_facility_verified", YES_NO
    )
    add_dropdown(
        form, form_headers_actual, "physical_vehicle_access_confirmed", YES_NO
    )
    add_dropdown(form, form_headers_actual, "vehicle_type_assessed", VEHICLE_OPTIONS)
    add_dropdown(form, form_headers_actual, "road_surface_observed", ROAD_SURFACE_OPTIONS)
    add_dropdown(form, form_headers_actual, "seasonal_access_concern", YES_NO)
    add_dropdown(form, form_headers_actual, "safe_vehicle_turnaround", YES_NO)
    add_dropdown(
        form, form_headers_actual, "nearest_road_connection_verified", YES_NO
    )
    add_dropdown(form, form_headers_actual, "follow_up_required", YES_NO)
    add_dropdown(form, form_headers_actual, "review_decision", DECISION_OPTIONS)
    add_dropdown(
        form, form_headers_actual, "routing_approval_recommendation",
        '"NO,UNDER_REVIEW,APPROVED"'
    )
    add_dropdown(form, form_headers_actual, "routing_approval_granted", '"NO"')

    for field in OBSERVATION_FIELDS:
        col = form_headers_actual.index(field) + 1
        for row_num in range(2, form.max_row + 1):
            form.cell(row_num, col).fill = PatternFill(
                "solid", fgColor="FFF2CC"
            )

    # Create a standalone CSV template with blank field-observation columns.
    csv_headers = source_columns + OBSERVATION_FIELDS
    with CSV_OUT.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=csv_headers)
        writer.writeheader()
        for row in rows:
            out_row = {col: row.get(col, "") for col in source_columns}
            out_row.update({field: "" for field in OBSERVATION_FIELDS})
            out_row["routing_approval_granted"] = "NO"
            writer.writerow(out_row)

    wb.save(WORKBOOK_OUT)

    # Reopen workbook and validate its key sheets and row counts.
    check = load_workbook(WORKBOOK_OUT, read_only=True, data_only=True)
    expected_sheet_counts = {
        "All Candidates": 256,
        "High Risk First": 26,
        "Field Observation Form": 256,
    }
    actual_sheet_counts = {
        name: check[name].max_row for name in expected_sheet_counts
    }
    check.close()

    if actual_sheet_counts != expected_sheet_counts:
        stop(f"Workbook row-count validation failed: {actual_sheet_counts}")

    audit_lines = [
        "PNIS PHASE 2G.19 - FIELDWORK EXECUTION PACKAGE",
        "=" * 72,
        f"Generated: {datetime.now().isoformat(timespec='seconds')}",
        "Status: COMPLETED - INPUT AND OUTPUT VALIDATION PASSED",
        "",
        "INPUT",
        str(INPUT),
        "",
        "INPUT VALIDATION",
        f"Input rows: {len(rows)}",
        f"Unique candidate IDs: {len({r['candidate_record_id'] for r in rows})}",
        "Duplicate candidate IDs: 0",
        "Priority order preserved: YES",
        "",
        "RISK COUNTS",
        f"HIGH_REVIEW: {len(high_risk)}",
        f"ELEVATED_REVIEW: {len(elevated)}",
        f"STANDARD_REVIEW: {len(standard)}",
        "",
        "OUTPUT VALIDATION",
        f"All Candidates sheet rows including header: {actual_sheet_counts['All Candidates']}",
        f"High Risk First sheet rows including header: {actual_sheet_counts['High Risk First']}",
        f"Field Observation Form rows including header: {actual_sheet_counts['Field Observation Form']}",
        "",
        "GUARDRAILS",
        "- Source worklist was read only.",
        "- Field observation values were left blank.",
        "- No routing approval was granted.",
        "- routing_approval_granted defaults to NO.",
        "- Existing source, workbook, audits, and routing graph were not edited.",
        "",
        "OUTPUT FILES",
        str(WORKBOOK_OUT),
        str(CSV_OUT),
        str(AUDIT_OUT),
    ]
    AUDIT_OUT.write_text("\n".join(audit_lines) + "\n", encoding="utf-8")

    print(f"Candidates: {len(rows)}")
    print(f"HIGH_REVIEW: {len(high_risk)}")
    print(f"ELEVATED_REVIEW: {len(elevated)}")
    print(f"STANDARD_REVIEW: {len(standard)}")
    print("Priority order preserved: YES")
    print("Routing approval granted: NO")
    print("\nOUTPUT FILES")
    print(WORKBOOK_OUT)
    print(CSV_OUT)
    print(AUDIT_OUT)
    print("\nSTATUS: COMPLETED - INPUT AND OUTPUT VALIDATION PASSED")


if __name__ == "__main__":
    main()
